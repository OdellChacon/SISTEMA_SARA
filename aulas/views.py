from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
import csv
from .models import Aula
from .forms import AulaForm  # Asegúrate de tener un formulario para Aula
from django.core.paginator import Paginator
import json

def aulas_list(request):
    aulas = Aula.objects.all()
    paginator = Paginator(aulas, 5)  # Mostrar 5 registros por página
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    return render(request, 'aulas/aulas_list.html', {'page_obj': page_obj, 'aulas': page_obj.object_list})

def registrar_aula(request):
    if request.method == 'POST':
        form = AulaForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('aulas_list')
    else:
        form = AulaForm()
    return render(request, 'aulas/registrar_aula.html', {'form': form})

def editar_aula(request, aula_id):
    aula = get_object_or_404(Aula, id=aula_id)
    if request.method == 'POST':
        if 'eliminar' in request.POST:  # Manejar eliminación
            aula.delete()
            return JsonResponse({'success': True})
        form = AulaForm(request.POST, instance=aula)
        if form.is_valid():
            form.save()
            return redirect('aulas_list')
    else:
        form = AulaForm(instance=aula)
    return render(request, 'aulas/editar_aula.html', {'form': form, 'aula': aula})

@csrf_exempt
def eliminar_aula(request, aula_id):
    if request.method == 'POST':
        aula = get_object_or_404(Aula, id=aula_id)
        aula.delete()
        return JsonResponse({'success': True, 'message': 'Aula eliminada correctamente.'})
    return JsonResponse({'success': False, 'message': 'Método no permitido.'}, status=405)

@csrf_exempt
def importar_aulas(request):
    if request.method == 'POST' and request.FILES.get('file'):
        file = request.FILES['file']
        filename = file.name.lower()
        def safe_value(val, dash=False):
            if val is None or str(val).strip() == '' or str(val).lower() == 'null':
                return '-' if dash else ''
            return str(val)
        if filename.endswith('.csv'):
            reader = csv.DictReader(file.read().decode('utf-8').splitlines())
            for row in reader:
                Aula.objects.create(
                    codigo_aula=safe_value(row.get('codigo_aula'), dash=True),
                    descripcion=safe_value(row.get('descripcion')),
                    capacidad=int(row.get('capacidad') or 0),
                    estatus=safe_value(row.get('estatus'), dash=True),
                    sede=safe_value(row.get('sede'), dash=True),
                    serial=safe_value(row.get('serial'))
                )
            return JsonResponse({'success': True})
        elif filename.endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            idx = {h: i for i, h in enumerate(headers)}
            for row in ws.iter_rows(min_row=2, values_only=True):
                def get(h, dash=False):
                    v = row[idx[h]] if h in idx and idx[h] < len(row) else None
                    return safe_value(v, dash)
                Aula.objects.create(
                    codigo_aula=get('codigo_aula', dash=True),
                    descripcion=get('descripcion'),
                    capacidad=int(get('capacidad') or 0),
                    estatus=get('estatus', dash=True),
                    sede=get('sede', dash=True),
                    serial=get('serial')
                )
            return JsonResponse({'success': True})
        else:
            return JsonResponse({'success': False, 'message': 'Formato de archivo no soportado'}, status=400)
    return JsonResponse({'success': False, 'message': 'Archivo no válido'}, status=400)

def exportar_aulas(request, format, scope):
    aulas = Aula.objects.all() if scope == "all" else Aula.objects.filter(id__in=scope.split(","))
    
    if format == "csv":
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="aulas.csv"'
        writer = csv.writer(response)
        writer.writerow(['codigo_aula', 'descripcion', 'capacidad', 'estatus', 'sede', 'serial'])
        for aula in aulas:
            writer.writerow([
                aula.codigo_aula,
                aula.descripcion,
                aula.capacidad,
                aula.estatus,
                aula.sede,
                aula.serial
            ])
        return response

    elif format == "excel":
        import openpyxl
        from openpyxl.utils import get_column_letter
        from openpyxl.writer.excel import save_virtual_workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['codigo_aula', 'descripcion', 'capacidad', 'estatus', 'sede', 'serial'])
        for aula in aulas:
            ws.append([
                aula.codigo_aula,
                aula.descripcion,
                aula.capacidad,
                aula.estatus,
                aula.sede,
                aula.serial
            ])
        response = HttpResponse(
            save_virtual_workbook(wb),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="aulas.xlsx"'
        return response

    return JsonResponse({'success': False, 'message': 'Formato no soportado'}, status=400)

@csrf_exempt
def eliminar_seleccionados(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        ids = data.get('ids', [])
        Aula.objects.filter(id__in=ids).delete()
        return JsonResponse({'success': True, 'message': 'Registros eliminados correctamente.'})
    return JsonResponse({'success': False, 'message': 'Método no permitido.'}, status=405)

def obtener_todos_los_ids(request):
    if request.method == 'GET':
        ids = list(Aula.objects.values_list('id', flat=True))
        return JsonResponse({'ids': ids})
    return JsonResponse({'error': 'Método no permitido'}, status=405)

def buscar_aulas(request):
    query = request.GET.get('q', '').strip()
    if query:
        aulas = Aula.objects.filter(
            codigo_aula__icontains=query
        ) | Aula.objects.filter(
            descripcion__icontains=query
        ) | Aula.objects.filter(
            estatus__icontains=query
        ) | Aula.objects.filter(
            sede__icontains=query
        ) | Aula.objects.filter(
            serial__icontains=query
        )
        results = [
            {
                'id': aula.id,
                'codigo_aula': aula.codigo_aula,
                'descripcion': aula.descripcion,
                'capacidad': aula.capacidad,
                'estatus': aula.estatus,
                'sede': aula.sede,
                'serial': aula.serial
            } for aula in aulas
        ]
    else:
        results = []
    return JsonResponse({'results': results})
